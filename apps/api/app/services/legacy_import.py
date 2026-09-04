import csv
import io
import json
import uuid
from datetime import date, datetime
from types import SimpleNamespace
from typing import Iterable
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.models.core import Branch
from app.models.patient import District, ImportBatch, ImportStaging, Locality, Patient, Region
from app.services.patients import create_patient, current_fiscal_year, search_patient_matches

# Financial data is explicitly not migrated. These columns are stripped before staging.
BANNED_FINANCE_TOKENS = {
    "bill", "billing", "amount", "payment", "paid", "due", "dues", "outstanding",
    "balance", "refund", "discount", "credit_note", "creditnote", "gst", "tax",
    "revenue", "package_amount", "package_value", "invoice", "receipt",
}

ALIASES = {
    "full_name": ["full_name", "name", "patient_name", "patient"],
    "mobile": ["mobile", "phone", "phone_number", "contact", "contact_number"],
    "alternate_mobile": ["alternate_mobile", "alt_mobile", "phone2", "alternate_phone"],
    "gender": ["gender", "sex"],
    "dob": ["dob", "date_of_birth", "birth_date"],
    "age": ["age", "years"],
    "region": ["region"],
    "district": ["district"],
    "locality": ["locality", "area"],
    "address_line": ["address", "address_line"],
    "legacy_patient_id": ["legacy_patient_id", "old_patient_id", "patient_id", "old_id", "id"],
}


def _norm_header(value) -> str:
    return "_".join(str(value or "").strip().lower().replace("-", " ").split())


def _is_financial_header(header: str) -> bool:
    h = _norm_header(header)
    return any(token == h or token in h for token in BANNED_FINANCE_TOKENS)


def _rows_from_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    return [dict(row) for row in csv.DictReader(io.StringIO(text))]


def _rows_from_xlsx(data: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    headers = [str(x or "").strip() for x in next(values, [])]
    return [dict(zip(headers, row)) for row in values]


def read_uploaded_rows(filename: str, data: bytes) -> list[dict]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return _rows_from_csv(data)
    if lower.endswith(".xlsx"):
        return _rows_from_xlsx(data)
    raise ValueError("Only CSV and XLSX patient imports are supported")


def detect_mapping(headers: Iterable[str], override: dict | None = None) -> dict[str, str]:
    normalized = {_norm_header(h): h for h in headers}
    mapping: dict[str, str] = {}
    override = override or {}
    for target, source in override.items():
        if source in headers:
            mapping[target] = source
    for target, aliases in ALIASES.items():
        if target in mapping:
            continue
        for alias in aliases:
            if _norm_header(alias) in normalized:
                mapping[target] = normalized[_norm_header(alias)]
                break
    return mapping


def _clean_value(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    return s if s else None


def _parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    s = str(value).strip()
    for fmt in ["%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"]:
        try: return datetime.strptime(s, fmt).date()
        except ValueError: pass
    return None


def _patient_payload_from_raw(db: Session, raw: dict, *, mapping: dict, organization_id: uuid.UUID, branch_id: uuid.UUID):
    def get(field):
        source = mapping.get(field)
        return _clean_value(raw.get(source)) if source else None

    full_name = get("full_name")
    if not full_name:
        return None, "Name is required"
    dob = _parse_date(raw.get(mapping.get("dob"))) if mapping.get("dob") else None
    dob_is_estimated = False
    age = get("age")
    if not dob and age:
        try:
            years = int(float(age))
            dob = date(datetime.now().year - years, 1, 1)
            dob_is_estimated = True
        except Exception:
            pass

    region_id = district_id = locality_id = None
    region_name, district_name, locality_name = get("region"), get("district"), get("locality")
    if region_name:
        region = db.scalar(select(Region).where(Region.name.ilike(region_name)).limit(1)); region_id = region.region_id if region else None
    if district_name:
        q = select(District).where(District.name.ilike(district_name))
        if region_id: q = q.where(District.region_id == region_id)
        district = db.scalar(q.limit(1)); district_id = district.district_id if district else None
        if district and not region_id: region_id = district.region_id
    if locality_name:
        q = select(Locality).where(Locality.locality_name.ilike(locality_name), Locality.approval_status == "APPROVED")
        if district_id: q = q.where(Locality.district_id == district_id)
        locality = db.scalar(q.limit(1)); locality_id = locality.locality_id if locality else None
        if locality and not district_id:
            district_id = locality.district_id
            district = db.get(District, district_id)
            region_id = district.region_id if district else region_id

    gender = get("gender")
    profile_status = "COMPLETE" if gender and dob and locality_id else "LEGACY_INCOMPLETE"
    payload = SimpleNamespace(
        organization_id=organization_id,
        full_name=full_name,
        mobile=get("mobile"), alternate_mobile=get("alternate_mobile"), gender=gender,
        dob=dob, dob_is_estimated=dob_is_estimated,
        registration_branch_id=branch_id,
        legacy_patient_id=get("legacy_patient_id"), profile_status=profile_status,
        region_id=region_id, district_id=district_id, locality_id=locality_id,
        address_line=get("address_line"), duplicate_override_confirmed=False,
    )
    return payload, None


def stage_import(db: Session, *, batch: ImportBatch, rows: list[dict], mapping: dict, organization_id: uuid.UUID, branch_id: uuid.UUID) -> dict:
    ignored = sorted({h for row in rows for h in row.keys() if _is_financial_header(h)})
    clean_headers = [h for h in (rows[0].keys() if rows else []) if not _is_financial_header(h)]
    effective_mapping = {k:v for k,v in mapping.items() if v in clean_headers}
    duplicate_count = error_count = 0
    for idx, source in enumerate(rows, start=2):
        safe_raw = {k:_clean_value(v) for k,v in source.items() if not _is_financial_header(k)}
        payload, error = _patient_payload_from_raw(db, safe_raw, mapping=effective_mapping, organization_id=organization_id, branch_id=branch_id)
        status, action, matched = "READY", "CREATE", None
        candidate_summaries = []
        if error:
            status, action, error_count = "ERROR", "SKIP", error_count + 1
        else:
            candidates = search_patient_matches(db, organization_id=organization_id, name=payload.full_name, mobile=payload.mobile, dob=payload.dob, gender=payload.gender, limit=5)
            candidate_summaries = [{
                "patient_id": str(x["patient_id"]), "uhid": x["uhid"], "full_name": x["full_name"],
                "age": x.get("age"), "gender": x.get("gender"), "primary_mobile": x.get("primary_mobile"),
                "match_reasons": x.get("match_reasons", []),
            } for x in candidates]
            if candidates:
                duplicate_count += 1; status = "DUPLICATE_REVIEW"; action = "REVIEW"; matched = candidates[0]["patient_id"]
        db.add(ImportStaging(import_batch_id=batch.import_batch_id, source_row_number=idx, raw_data={"data":safe_raw,"mapping":effective_mapping,"candidates":candidate_summaries}, validation_status=status, matched_patient_id=matched, import_action=action, error_message=error))
    batch.total_rows = len(rows); batch.error_rows = error_count; batch.duplicate_rows = duplicate_count; batch.status = "STAGED"
    return {"ignored_financial_columns": ignored, "mapping": effective_mapping}


def commit_import(db: Session, *, batch: ImportBatch, organization_id: uuid.UUID, branch_id: uuid.UUID, actor_user_id: uuid.UUID):
    branch = db.get(Branch, branch_id)
    if not branch or branch.organization_id != organization_id:
        raise ValueError("Registration branch not found")
    rows = list(db.scalars(select(ImportStaging).where(ImportStaging.import_batch_id == batch.import_batch_id).order_by(ImportStaging.source_row_number).with_for_update()).all())
    created = linked = skipped = errors = 0
    for row in rows:
        if row.import_action == "LINK_EXISTING":
            if not row.matched_patient_id or not db.get(Patient, row.matched_patient_id):
                row.validation_status = "ERROR"; row.error_message = "Matched Patient is missing"; errors += 1
            else:
                row.validation_status = "IMPORTED_LINKED"; linked += 1
            continue
        if row.import_action not in {"CREATE", "CREATE_SEPARATE"}:
            skipped += 1; continue
        payload, error = _patient_payload_from_raw(db, row.raw_data.get("data", {}), mapping=row.raw_data.get("mapping", {}), organization_id=organization_id, branch_id=branch_id)
        if error:
            row.validation_status = "ERROR"; row.error_message = error; errors += 1; continue
        if row.import_action == "CREATE_SEPARATE": payload.duplicate_override_confirmed = True
        try:
            p = create_patient(db, payload=payload, actor_user_id=actor_user_id, fiscal_year=current_fiscal_year(), branch_code=branch.branch_code)
            row.matched_patient_id = p.patient_id; row.validation_status = "IMPORTED"; created += 1
        except Exception as exc:
            row.validation_status = "ERROR"; row.error_message = str(exc)[:1000]; errors += 1
    total_errors = sum(1 for r in rows if r.validation_status == "ERROR")
    batch.imported_rows = created + linked; batch.error_rows = total_errors; batch.status = "COMPLETED_WITH_ERRORS" if total_errors else "COMPLETED"
    return {"created":created,"linked":linked,"skipped":skipped,"errors":errors}
